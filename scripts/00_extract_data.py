"""
Data extraction script — pulls the cleaned CSVs used by the four training scripts
out of the three raw CEPPI Excel workbooks. Run this first.

Outputs (into ./data/):
  fo_flux_shortterm.csv            Fig. 3A short-term FO flux (dataset A)
  sample_analysis.csv              Feed-in/out, draw-out lab analyses (dataset A)
  ro_nacl_calibration.csv          RO rejection/flux vs TMP, 30/40/50 bar (water_tests)
  fo_flux_water_baseline.csv       FO with DI water feed, Fig. 2A baseline (water_tests)
  fo_flux_longterm_rep1_full.csv   Long-term replicate 1, full sensor features (dataset B)
  fo_flux_longterm_rep1.csv        Long-term replicate 1, time+Jw only, as authors computed it
"""
import openpyxl
import pandas as pd
import numpy as np
import os

os.makedirs("data", exist_ok=True)

UPLOADS = "data"  # adjust path if running elsewhere

# ---------------------------------------------------------------------------
# 1. FO permeate flux, short-term real-wastewater run (dataset A) -> Fig. 3A
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(f"{UPLOADS}/CEPPI_2025_dataset_A.xlsx", data_only=True)
ws = wb["FO permeate flux"]
rows = list(ws.iter_rows(min_row=3, values_only=True))
recs = []
for r in rows:
    t, jw, sd = r[3], r[4], r[5]          # columns D,E,F hold Time / Jw / SD
    if t is not None and jw is not None:
        recs.append((t, jw, sd))
fo_short = pd.DataFrame(recs, columns=["time_min", "Jw", "SD"]).dropna(subset=["Jw"])
fo_short.to_csv("data/fo_flux_shortterm.csv", index=False)
print("fo_flux_shortterm:", fo_short.shape)

# ---------------------------------------------------------------------------
# 2. Sample analysis: COD/TOC/TDS/conductivity for feed-in/out & draw-out
#    at each water-recovery level (dataset A) -> Table 4 / Table 5
# ---------------------------------------------------------------------------
ws = wb["Sample analysis"]
rows = list(ws.iter_rows(min_row=2, max_row=25, values_only=True))
recs = []
for r in rows:
    sid, stype, pH, cond, osm, tds = r[0], r[1], r[2], r[3], r[4], r[5]
    toc, tn = r[10], r[11]
    if stype:
        recs.append((sid, stype, pH, cond, osm, tds, toc, tn))
sample = pd.DataFrame(
    recs,
    columns=["sample_id", "sample_type", "pH", "conductivity_uS_cm",
             "osmolality", "TDS_mg_L", "TOC_mg_L", "TN_mg_L"],
)
# sample_type looks like "Feed in 25%" / "Feed out 90%" / "Draw out 25%"
sample[["stream", "recovery_pct"]] = sample["sample_type"].str.extract(
    r"(Feed in|Feed out|Draw out)\s+(\d+)%"
)
sample["recovery_pct"] = pd.to_numeric(sample["recovery_pct"])
sample = sample.dropna(subset=["stream"])
sample.to_csv("data/sample_analysis.csv", index=False)
print("sample_analysis:", sample.shape)

# ---------------------------------------------------------------------------
# 3. RO quality-control test: 0.5 M NaCl feed at 30 / 40 / 50 bar
#    (water_tests workbook) -> Fig. 2B
# ---------------------------------------------------------------------------
wb2 = openpyxl.load_workbook(f"{UPLOADS}/CEPPI_2025_water_tests.xlsx", data_only=True)
ws = wb2["RO - NaCl 0.5 M"]
recs = []
cur = None
for row in ws.iter_rows(min_row=12, max_row=25, values_only=True):
    test, flow, tmp, perm_cond, feed_cond, rej, perm_flow, rec = (
        row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8]
    )
    if isinstance(test, (int, float)):
        if cur:
            recs.append(cur)
        cur = {
            "test": test, "flow_Lh": flow, "TMP_bar": tmp,
            "perm_cond_list": [perm_cond], "feed_cond_list": [feed_cond],
            "rejection": rej, "perm_flow_Lh": perm_flow, "recovery": rec,
        }
    elif test is None and perm_cond is not None and cur is not None and str(row[1]) != "AVERAGE":
        cur["perm_cond_list"].append(perm_cond)
        cur["feed_cond_list"].append(feed_cond)
if cur:
    recs.append(cur)
for c in recs:
    c["perm_cond_avg_mScm"] = sum(c["perm_cond_list"]) / len(c["perm_cond_list"])
    c["feed_cond_avg_mScm"] = sum(c["feed_cond_list"]) / len(c["feed_cond_list"])
    del c["perm_cond_list"], c["feed_cond_list"]
ro_cal = pd.DataFrame(recs)
ro_cal.to_csv("data/ro_nacl_calibration.csv", index=False)
print("ro_nacl_calibration:", ro_cal.shape)

# ---------------------------------------------------------------------------
# 4. FO quality-control test: DI water feed, 0.5 M NaCl draw (water_tests) -> Fig. 2A
# ---------------------------------------------------------------------------
ws = wb2["FO - water as feed"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
recs = []
for r in rows:
    t, jw, js, jsjw = r[0], r[1], r[2], r[3]
    if t is not None and jw is not None:
        recs.append((t, jw, js, jsjw))
fo_base = pd.DataFrame(recs, columns=["time_min", "Jw", "Js", "Js_Jw"])
fo_base.to_csv("data/fo_flux_water_baseline.csv", index=False)
print("fo_flux_water_baseline:", fo_base.shape)

# ---------------------------------------------------------------------------
# 5. Longer-term run, replicate 1, full sensor log (dataset B) -> Fig. 6A
#    Jw is recomputed from raw flow channels using the SAME formula the
#    original workbook uses: (feed-in flow - feed-out flow) / membrane area
# ---------------------------------------------------------------------------
wb3 = openpyxl.load_workbook(f"{UPLOADS}/CEPPI_2025_dataset_B.xlsx", data_only=True)
ws = wb3["Raw replicate 1"]
header = [c.value for c in ws[1]]
rows = list(ws.iter_rows(min_row=2, values_only=True))
df = pd.DataFrame(rows, columns=header)

cols = ["Experiment Time (min)", "F2-FO-Feed (l/hr)", "F3-FO-Feed-Out (l/hr)",
        "C1-FO-Draw-Cond. (mS/cm)", "C2-FO-Feed-Cond. (mS/cm)",
        "C1-FO-Draw-Temp. (C)", "C2-FO-Feed-Temp. (C)",
        "RO-TMP (barg)", "FO-TMP-real (bar)"]
d = df[cols].apply(pd.to_numeric, errors="coerce").dropna().reset_index(drop=True)
d = d.rename(columns={
    "Experiment Time (min)": "time_min",
    "F2-FO-Feed (l/hr)": "feed_in_Lh",
    "F3-FO-Feed-Out (l/hr)": "feed_out_Lh",
    "C1-FO-Draw-Cond. (mS/cm)": "draw_cond_mScm",
    "C2-FO-Feed-Cond. (mS/cm)": "feed_cond_mScm",
    "C1-FO-Draw-Temp. (C)": "draw_temp_C",
    "C2-FO-Feed-Temp. (C)": "feed_temp_C",
    "RO-TMP (barg)": "RO_TMP_bar",
    "FO-TMP-real (bar)": "FO_TMP_bar",
})
AREA = 13.8  # m^2, HFFO14 membrane active area (Set points sheet)
d["Jw"] = (d["feed_in_Lh"] - d["feed_out_Lh"]) / AREA
d = d[(d["Jw"] > 0) & (d["Jw"] < 15)]  # drop sensor glitches / tank-purge artifacts
d.to_csv("data/fo_flux_longterm_rep1_full.csv", index=False)
print("fo_flux_longterm_rep1_full:", d.shape)

# Also save the authors' own pre-computed Jw column (Calculations FO sheet, Day 2),
# useful as a sanity check against the recomputation above.
ws = wb3["Calculations FO"]
rows = list(ws.iter_rows(min_row=3, values_only=True))
recs = []
for r in rows:
    t, jw2 = r[0], r[1]
    if isinstance(t, (int, float)) and isinstance(jw2, (int, float)):
        recs.append((t, jw2))
fo_long_r1 = pd.DataFrame(recs, columns=["time_min", "Jw"])
fo_long_r1.to_csv("data/fo_flux_longterm_rep1.csv", index=False)
print("fo_flux_longterm_rep1 (authors' formula):", fo_long_r1.shape)

print("\nNote: replicate 2's flux is NOT extracted here. The workbook's own formula for it\n"
      "references a deleted column (#REF! error), and it lacks a feed-outlet-flow sensor\n"
      "channel, so it cannot be reconstructed with the same method as replicate 1.")
